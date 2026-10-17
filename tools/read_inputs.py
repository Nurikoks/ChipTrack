"""Dump case document text and raw label sheet (Phase 0 audit helper)."""
import sys
import docx
import openpyxl

doc_path, xlsx_path = sys.argv[1], sys.argv[2]

d = docx.Document(doc_path)
print("=== CASE DOCUMENT ===")
for p in d.paragraphs:
    if p.text.strip():
        print(p.text)
for t in d.tables:
    print("--- table ---")
    for r in t.rows:
        print(" | ".join(c.text.strip() for c in r.cells))

print("\n=== LABELS ===")
wb = openpyxl.load_workbook(xlsx_path)
for ws in wb.worksheets:
    print("sheet:", ws.title)
    for row in ws.iter_rows(values_only=True):
        print(repr(row))
